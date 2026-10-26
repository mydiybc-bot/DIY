#!/usr/bin/env python3
"""104 匯入範本（時間排班）交付前檢查。

正本：~/.claude/skills/diybc-scheduling/scripts/check_104.py（門市排班工作台 schedule_ws.py 直接 import 這份副本；改正本要同步複製過來）。

用法：
  python3 check_104.py 已排班.xlsx --template 店長上傳的原始範本.xlsx
  可加：--max-consec 5            該店連上上限（預設 5；超過 6 天一律算錯誤）
        --expect 0021=168,0030=176  報告裡每人的時數，逐人比對

合法的出勤類型、標籤、崗位、時間格式一律讀範本自己的 hidden 分頁，不寫死。
結果：❌＝會被 104 擋下或違反勞基法，一定要改；⚠＝請店長確認。
有任何 ❌ 時結束碼為 1。
"""
import argparse
import datetime as dt
import re
import sys
import warnings

import openpyxl

warnings.filterwarnings("ignore")

FIELDS = ["出勤類型", "起", "迄", "崗位", "標籤", "當日標籤"]
REST_TYPES = {"例假日", "休息日", "國定假日"}
OK_TYPES = REST_TYPES | {"特休"}
OTHER_STORE_TAGS = {"支援", "台北內訓課", "區會", "經管會", "交流會議"}
# 104 班表名稱關鍵字 → 崗位簡稱（判斷「崗位填成本店」用）
STORE_POS = [("高雄", "高雄店"), ("精明", "精明店"), ("草悟道", "草悟道店"), ("台南", "台南店"),
             ("新竹", "新竹店"), ("南京", "南京店"), ("士林", "士林店"), ("新店", "新店店"),
             ("板橋", "板橋店"), ("A13", "A13店"), ("中壢", "中壢店"), ("藝文", "藝文店")]


class Report:
    def __init__(self):
        self.err, self.warn = [], []

    def e(self, msg):
        self.err.append(msg)

    def w(self, msg):
        self.warn.append(msg)


def s(v):
    """儲存格值 → 字串（空白回 ''）。"""
    return "" if v is None else str(v)


def mins(t):
    """'10:20' → 620；'次日01:00' → 1500。"""
    nxt = t.startswith("次日")
    h, m = t.replace("次日", "").split(":")
    return int(h) * 60 + int(m) + (1440 if nxt else 0)


def hm(m):
    return f"{m // 60}h" if m % 60 == 0 else f"{m / 60:.2f}h"


def read_lists(wb):
    """從範本 hidden 分頁讀合法清單（名稱定義 dayTypes／tagStrs／positionStrs／timeStrs／endTimeStrs）。"""
    out = {}
    for name in ["dayTypes", "tagStrs", "positionStrs", "timeStrs", "endTimeStrs"]:
        dn = wb.defined_names.get(name)
        vals = set()
        if dn is not None:
            for title, ref in dn.destinations:
                for row in wb[title][ref.replace("$", "")]:
                    for c in row:
                        if c.value not in (None, ""):
                            vals.add(str(c.value))
        out[name] = vals
    return out


def layout(ws):
    """找標題列、每天的欄位組、每位員工的 4 列。"""
    hdr = next(r for r in range(1, 30) if s(ws.cell(r, 1).value) == "員工編號")
    days, col = [], 4
    while col <= ws.max_column:
        label = s(ws.cell(hdr - 1, col).value)
        m = re.match(r"(\d{1,2})/(\d{1,2})\((.)\)", label)
        if not m:
            break
        days.append({"col": col, "label": label[:5], "wd": m.group(3)})
        col += 6
    emps, r = [], hdr + 1
    while r <= ws.max_row:
        if s(ws.cell(r, 3).value) == "第1段" and s(ws.cell(r, 1).value):
            emps.append({"row": r, "id": s(ws.cell(r, 1).value), "name": s(ws.cell(r, 2).value)})
            r += 4
        else:
            r += 1
    return hdr, days, emps


def used(ws):
    """有內容的最後一列、最後一欄。"""
    cells = [(c.row, c.column) for row in ws.iter_rows() for c in row if c.value not in (None, "")]
    return (max(r for r, _ in cells), max(c for _, c in cells)) if cells else (0, 0)


def compare_structure(rep, wb, tb):
    """跟店長上傳的原始範本比：除了排班格子，其他一律不能動。"""
    if wb.sheetnames != tb.sheetnames:
        rep.e(f"分頁名稱或順序被改了：範本 {tb.sheetnames} → 現在 {wb.sheetnames}")
        return
    for a, b in zip(wb.worksheets, tb.worksheets):
        if a.sheet_state != b.sheet_state:
            rep.e(f"分頁「{a.title}」顯示／隱藏狀態被改了")
    if sorted(wb.defined_names.keys()) != sorted(tb.defined_names.keys()):
        rep.e("下拉選單用的名稱定義（dayTypes 等）不見或被改了")
    for name in ["hidden", "班別明細"]:
        if name in tb.sheetnames:
            va = [[c.value for c in r] for r in wb[name].iter_rows()]
            vb = [[c.value for c in r] for r in tb[name].iter_rows()]
            if va != vb:
                rep.e(f"分頁「{name}」內容被改了（下拉選單清單），請用原始範本重填")
    ws, ts = wb.worksheets[0], tb.worksheets[0]
    # 只比有內容的範圍（程式重存常會丟掉最後的空白列，不影響匯入）
    if used(ws) != used(ts):
        rep.e(f"第一頁有內容的範圍被改了：範本 {used(ts)[0]} 列×{used(ts)[1]} 欄 → 現在 {used(ws)[0]}×{used(ws)[1]}（不可增刪列欄）")
    ma = {str(m) for m in ws.merged_cells.ranges}
    mb = {str(m) for m in ts.merged_cells.ranges}
    if ma != mb:
        rep.e(f"合併儲存格被改了：少了 {len(mb - ma)} 個、多了 {len(ma - mb)} 個")
    da = {(d.formula1, str(d.sqref)) for d in ws.data_validations.dataValidation}
    db = {(d.formula1, str(d.sqref)) for d in ts.data_validations.dataValidation}
    if da != db:
        rep.e(f"下拉選單（資料驗證）掉了或變了：範本 {len(db)} 組 → 現在 {len(da)} 組。"
              "常見原因是用程式整檔重寫；請改成只改儲存格值或直接改工作表 XML")
    hdr, days, emps = layout(ts)
    data_cols = {d["col"] + i for d in days for i in range(6)}
    data_rows = {e["row"] + i for e in emps for i in range(4)}
    changed = 0
    for r in range(1, ts.max_row + 1):
        for c in range(1, ts.max_column + 1):
            if r in data_rows and c in data_cols:
                continue
            if s(ws.cell(r, c).value) != s(ts.cell(r, c).value):
                changed += 1
                if changed <= 5:
                    rep.e(f"排班格子以外的內容被改了：{ws.cell(r, c).coordinate}「{s(ts.cell(r, c).value)}」→「{s(ws.cell(r, c).value)}」")
    if changed > 5:
        rep.e(f"……排班格子以外共 {changed} 格被改（標題、工號、姓名都不能動）")


def check(path, template, max_consec, expect):
    rep = Report()
    wb = openpyxl.load_workbook(path)
    ws = wb.worksheets[0]
    if template:
        compare_structure(rep, wb, openpyxl.load_workbook(template))
    else:
        rep.w("沒有提供原始範本（--template），格式完整性只做基本檢查")
        if not ws.data_validations.dataValidation:
            rep.e("第一頁沒有任何下拉選單（資料驗證），格式很可能被弄壞")
    L = read_lists(wb if not template else openpyxl.load_workbook(template))
    if not L["dayTypes"] or not L["timeStrs"]:
        rep.e("讀不到範本的合法清單（hidden 分頁），無法檢查內容；請用店長當月下載的原始範本")
        return rep, [], ws.title
    own_pos = next((p for k, p in STORE_POS if k in ws.title), None)
    hdr, days, emps = layout(ws)
    rows = []
    for emp in emps:
        who = f"{emp['name']}({emp['id']})"
        day_info = []  # 每天：(類型, [(起,迄)], 當日標籤)
        for d in days:
            c0, r0, where = d["col"], emp["row"], f"{who} {d['label']}"
            cell = lambda rr, i: ws.cell(rr, c0 + i)
            for rr in range(r0, r0 + 4):
                for i in range(6):
                    v = cell(rr, i).value
                    if isinstance(v, (dt.time, dt.datetime, int, float)) and v != "":
                        rep.e(f"{where} {cell(rr, i).coordinate}：存成 Excel 時間／數字格式（{v}），要存成文字，例如「10:20」")
            dtype = s(cell(r0, 0).value)
            dtag = s(cell(r0, 5).value)
            for label, val, allowed in [("出勤類型", dtype, L["dayTypes"]), ("當日標籤", dtag, L["tagStrs"])]:
                if val and val not in allowed:
                    hint = "（前後多了空白）" if val.strip() in allowed else ""
                    rep.e(f"{where} {label}「{val}」不在 104 清單{hint}；可用：{'、'.join(sorted(allowed))}")
            segs, seen_empty = [], False
            for k in range(4):
                rr = r0 + k
                st, en = s(cell(rr, 1).value), s(cell(rr, 2).value)
                pos, tag = s(cell(rr, 3).value), s(cell(rr, 4).value)
                seg = f"{where} 第{k + 1}段"
                if not st and not en:
                    if pos or tag:
                        rep.e(f"{seg}：有崗位／標籤但沒有工作時間（104 規定標籤、崗位要搭配一組時間）")
                    seen_empty = True
                    continue
                if seen_empty:
                    rep.w(f"{seg}：上面一段是空的，時段請由第1段往下連續填")
                if not st or not en:
                    rep.e(f"{seg}：起訖時間只填了一邊（起「{st}」迄「{en}」）")
                    continue
                bad = False
                for label, val, allowed in [("起", st, L["timeStrs"]), ("迄", en, L["endTimeStrs"])]:
                    if isinstance(cell(rr, 1 if label == "起" else 2).value, (dt.time, dt.datetime)):
                        bad = True  # 上面已報「存成 Excel 時間格式」
                    elif val not in allowed:
                        bad = True
                        hint = "（前後多了空白）" if val.strip() in allowed else "（格式要像 10:20，5 分鐘一格，半形冒號）"
                        rep.e(f"{seg}：{label}時間「{val}」104 不收{hint}")
                if pos and pos not in L["positionStrs"]:
                    hint = "（前後多了空白）" if pos.strip() in L["positionStrs"] else ""
                    rep.e(f"{seg}：崗位「{pos}」不在 104 清單{hint}；可用：{'、'.join(sorted(L['positionStrs']))}")
                tags = [t for t in tag.split("、")] if tag else []
                for t in tags:
                    if t not in L["tagStrs"]:
                        hint = "（前後多了空白）" if t.strip() in L["tagStrs"] else ""
                        rep.e(f"{seg}：標籤「{t}」不在 104 清單{hint}；可用：{'、'.join(sorted(L['tagStrs']))}")
                if pos and pos == own_pos:
                    rep.w(f"{seg}：崗位填了本店「{pos}」，在本店上班的時段崗位要留空")
                elif pos and not (set(tags) & OTHER_STORE_TAGS):
                    rep.w(f"{seg}：崗位填了「{pos}」但沒打支援／台北內訓課／區會等標籤，請確認")
                if bad:
                    continue
                a, b = mins(st), mins(en)
                if b <= a:
                    rep.e(f"{seg}：迄 {en} 沒有晚於起 {st}")
                    continue
                if segs and a < segs[-1][1]:
                    rep.e(f"{seg}：{st}–{en} 跟上一段重疊或順序顛倒")
                segs.append((a, b))
            if dtype == "例假日" and segs:
                rep.e(f"{where}：例假日不能排上班；要上班請改「休息日」並排時段，或把例假換到別天")
            if dtype and dtype not in OK_TYPES:
                rep.w(f"{where}：出勤類型「{dtype}」，排班流程只用例假日／休息日／國定假日／特休，請確認")
            if dtype == "特休" and dtag != "特休":
                rep.w(f"{where}：出勤類型是特休，當日標籤也要打「特休」")
            day_info.append((dtype, segs, dtag))
        rows.append((emp, who, day_info))
    rule_checks(rep, days, rows, max_consec, expect)
    return rep, rows, ws.title


def rule_checks(rep, days, rows, max_consec, expect):
    for emp, who, info in rows:
        fulltime = any(t == "例假日" for t, _, _ in info)
        run, longest_run, run_start = 0, 0, None
        for i, (t, segs, _) in enumerate(info):
            day = days[i]["label"]
            total = sum(b - a for a, b in segs)
            if total > 12 * 60:
                rep.e(f"{who} {day}：當天 {hm(total)}，超過 12 小時上限")
            if segs:
                run += 1
                if run == 1:
                    run_start = day
                if run == 7:
                    rep.e(f"{who}：{run_start} 起連續上班超過 6 天（違反七休一）")
                elif run == max_consec + 1:
                    rep.w(f"{who}：{run_start} 起連上 {run} 天，超過本店上限 {max_consec} 天")
            else:
                run = 0
                if fulltime and not t:
                    rep.w(f"{who} {day}：正職這天沒排班也沒標休假類型（例假／休息／國定）")
            if i > 0 and segs and info[i - 1][1]:
                gap = segs[0][0] + 1440 - info[i - 1][1][-1][1]
                if gap < 11 * 60:
                    rep.w(f"{who} {days[i - 1]['label']}→{day}：下班到隔天上班只隔 {hm(gap)}，不到 11 小時")
        if fulltime:
            # 104 公司行事曆以週日為一週第一天；只檢查整週都在本月的週
            week = []
            for i, d in enumerate(days):
                if d["wd"] == "日":
                    week = [i]
                elif week:
                    week.append(i)
                if len(week) == 7:
                    types = [info[j][0] for j in week]
                    span = f"{days[week[0]]['label']}–{days[week[-1]]['label']}"
                    if types.count("例假日") != 1:
                        rep.w(f"{who} {span}：這週例假日 {types.count('例假日')} 天（應為 1 天）")
                    if "休息日" not in types:
                        rep.w(f"{who} {span}：這週沒有休息日")
                    week = []
        if emp["id"] in expect:
            got = sum(b - a for _, segs, _ in info for a, b in segs) / 60
            if abs(got - expect[emp["id"]]) > 0.01:
                rep.e(f"{who}：班表加總 {round(got, 2):g}h，跟報告的 {expect[emp['id']]:g}h 對不上")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("file")
    ap.add_argument("--template")
    ap.add_argument("--max-consec", type=int, default=5)
    ap.add_argument("--expect", default="")
    a = ap.parse_args()
    expect = {}
    for part in filter(None, a.expect.split(",")):
        k, v = part.split("=")
        expect[k.strip()] = float(v)
    rep, rows, title = check(a.file, a.template, a.max_consec, expect)

    print(f"=== 104 匯入檔檢查：{title} ===")
    if rows:
        print("\n【每人時數】（含特休保留時段、休息日／國定假日出勤）")
        print("工號\t姓名\t上班天\t總時數\t例假\t休息\t國定\t特休")
        for emp, _, info in rows:
            tot = sum(b - a for _, segs, _ in info for a, b in segs)
            work = sum(1 for _, segs, _ in info if segs)
            cnt = lambda x: sum(1 for t, _, _ in info if t == x)
            print(f"{emp['id']}\t{emp['name']}\t{work}\t{round(tot / 60, 2):g}h\t{cnt('例假日')}\t{cnt('休息日')}\t{cnt('國定假日')}\t{cnt('特休')}")
    print(f"\n【❌ 一定要改】{len(rep.err)} 項")
    for m in rep.err:
        print("❌ " + m)
    print(f"\n【⚠ 請店長確認】{len(rep.warn)} 項")
    for m in rep.warn:
        print("⚠ " + m)
    ok = not rep.err
    print("\n結果：" + ("✅ 通過，可以交給店長匯入" if ok else f"❌ 不通過（{len(rep.err)} 項要改），改完再跑一次"))
    sys.exit(0 if ok else 1)


if __name__ == "__main__":
    main()
